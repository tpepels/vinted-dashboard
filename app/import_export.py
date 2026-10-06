"""Generic CSV/TSV/XLSX inventory import/export.

The parser is connector-neutral.  It detects common column names, exposes a
preview before writes, refuses ambiguous/invalid rows, and applies confirmed
changes transactionally to the master InventoryItem table.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models
from app.constants import ItemCategory, ItemStatus, KNOWN_ITEM_CATEGORIES
from app.workspace_bootstrap import clean_isbn, normalize_sku


CANONICAL_FIELDS = (
    "sku",
    "title",
    "category",
    "quantity",
    "condition",
    "cost",
    "price",
    "currency",
    "location",
    "notes",
    "barcode",
    "author",
    "isbn",
    "subtitle",
    "publisher",
    "edition",
    "binding",
    "language",
    "publish_date",
    "publication_year",
    "pages",
    "brand",
    "size",
    "colour",
    "material",
    "measurements",
)

FIELD_ALIASES: dict[str, set[str]] = {
    "sku": {
        "sku", "stock number", "stock no", "inventory no", "inventory number",
        "book id", "bookid", "record number", "item number", "seller sku",
    },
    "title": {"title", "book name", "name", "item title", "product title"},
    "category": {"category", "type", "item type"},
    "quantity": {"quantity", "qty", "stock", "available"},
    "condition": {"condition", "item condition"},
    "cost": {"cost", "cost price", "purchase price", "buy price"},
    "price": {"price", "sell price", "selling price", "asking price"},
    "currency": {"currency", "currency code"},
    "location": {"location", "storage", "storage location", "shelf", "bin"},
    "notes": {"notes", "note"},
    "barcode": {"barcode", "gtin", "ean", "upc", "global unique id"},
    "author": {"author", "writer"},
    "isbn": {"isbn", "isbn10", "isbn13", "isbn-10", "isbn-13"},
    "subtitle": {"subtitle", "sub title"},
    "publisher": {"publisher"},
    "edition": {"edition"},
    "binding": {"binding", "format", "physical format"},
    "language": {"language", "lang"},
    "publish_date": {"publish date", "publication date", "date published"},
    "publication_year": {"publication year", "year", "published"},
    "pages": {"pages", "page count", "number of pages"},
    "brand": {"brand", "make"},
    "size": {"size"},
    "colour": {"colour", "color"},
    "material": {"material", "fabric"},
    "measurements": {"measurements", "measurement"},
}

ATTRIBUTE_FIELDS = {
    "barcode", "author", "isbn", "subtitle", "publisher", "edition",
    "binding", "language", "publish_date", "publication_year", "pages",
    "brand", "size", "colour", "material", "measurements",
}


@dataclass
class ParsedTable:
    file_type: str
    headers: list[str]
    rows: list[dict[str, Any]]


def _norm_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode the uploaded text file")


def _file_type(filename: str) -> str:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix in {"xlsx", "xlsm"}:
        return "xlsx"
    if suffix == "tsv":
        return "tsv"
    if suffix in {"csv", "txt"}:
        return "csv"
    raise ValueError("Supported file types are CSV, TSV and XLSX")


def parse_table(filename: str, content: bytes) -> ParsedTable:
    kind = _file_type(filename)
    if kind == "xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("The XLSX file could not be read") from exc
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        first = next(values, None)
        if not first:
            return ParsedTable(file_type="xlsx", headers=[], rows=[])
        headers = [str(value or "").strip() for value in first]
        rows: list[dict[str, Any]] = []
        for row in values:
            if not any(value not in (None, "") for value in row):
                continue
            rows.append({
                headers[index]: (row[index] if index < len(row) else None)
                for index in range(len(headers))
                if headers[index]
            })
        return ParsedTable(file_type="xlsx", headers=[h for h in headers if h], rows=rows)

    text = _decode_text(content)
    if not text.strip():
        return ParsedTable(file_type=kind, headers=[], rows=[])
    sample = text[:20000]
    delimiter = "\t" if kind == "tsv" else None
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    headers = [str(value or "").strip() for value in (reader.fieldnames or []) if value]
    rows = [
        {str(key or "").strip(): value for key, value in raw.items() if key}
        for raw in reader
        if any(value not in (None, "") for value in raw.values())
    ]
    return ParsedTable(
        file_type="tsv" if delimiter == "\t" else "csv",
        headers=headers,
        rows=rows,
    )


def suggest_mapping(headers: Iterable[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for header in headers:
        normalized = _norm_header(header)
        for field, aliases in FIELD_ALIASES.items():
            if field in used:
                continue
            if normalized in aliases:
                mapping[header] = field
                used.add(field)
                break
    return mapping


def _int(value: Any, default: int | None = None) -> int | None:
    if value in (None, ""):
        return default
    try:
        return int(float(str(value).strip().replace(",", ".")))
    except (TypeError, ValueError):
        return default


def _money(value: Any) -> int | None:
    if value in (None, ""):
        return None
    text = str(value).strip().replace("\u00a0", " ")
    match = re.search(r"-?[0-9]+(?:[.,][0-9]{1,2})?", text)
    if not match:
        return None
    try:
        return round(float(match.group(0).replace(",", ".")) * 100)
    except ValueError:
        return None


def _mapped(raw: dict[str, Any], mapping: dict[str, str]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for source, target in mapping.items():
        if target not in CANONICAL_FIELDS:
            continue
        value = raw.get(source)
        if value in (None, ""):
            continue
        result[target] = value
    return result


def _category(
    value: Any,
    mapped: dict[str, Any],
    default_category: str = ItemCategory.GENERAL,
) -> str:
    raw = str(value or "").strip().lower()
    if raw in KNOWN_ITEM_CATEGORIES:
        return raw
    if any(mapped.get(key) for key in ("isbn", "author", "publisher", "binding")):
        return ItemCategory.BOOK
    if any(mapped.get(key) for key in ("brand", "size", "colour", "material")):
        return ItemCategory.CLOTHING
    return (
        default_category
        if default_category in KNOWN_ITEM_CATEGORIES
        else ItemCategory.GENERAL
    )


def _row_payload(
    raw: dict[str, Any],
    mapping: dict[str, str],
    *,
    default_category: str = ItemCategory.GENERAL,
) -> tuple[dict[str, Any], list[str]]:
    mapped = _mapped(raw, mapping)
    errors: list[str] = []
    sku = normalize_sku(mapped.get("sku"))
    title = str(mapped.get("title") or "").strip()
    if not sku:
        errors.append("missing SKU")
    if not title:
        errors.append("missing title")

    attributes: dict[str, Any] = {}
    for key in ATTRIBUTE_FIELDS:
        value = mapped.get(key)
        if value in (None, ""):
            continue
        if key == "isbn":
            cleaned = clean_isbn(value)
            if cleaned:
                attributes[key] = cleaned
            else:
                errors.append("invalid ISBN")
        elif key == "barcode":
            cleaned = re.sub(r"[^0-9A-Za-z]", "", str(value))
            if cleaned:
                attributes[key] = cleaned
        elif key in {"publication_year", "pages"}:
            number = _int(value, None)
            if number is not None and number >= 0:
                attributes[key] = number
            else:
                errors.append(f"invalid {key.replace('_', ' ')}")
        else:
            attributes[key] = str(value).strip()

    quantity = _int(mapped.get("quantity"), 1)
    if quantity is None or quantity < 0:
        errors.append("invalid quantity")
        quantity = 1

    payload = {
        "sku": sku,
        "title": title,
        "category": _category(
            mapped.get("category"),
            mapped,
            default_category=default_category,
        ),
        "quantity": quantity,
        "condition": str(mapped.get("condition") or "").strip() or None,
        "cost_cents": _money(mapped.get("cost")),
        "currency": str(mapped.get("currency") or "EUR").strip().upper()[:3] or "EUR",
        "location": str(mapped.get("location") or "").strip() or None,
        "notes": str(mapped.get("notes") or "").strip() or None,
        "attributes": attributes,
    }
    sell_price = _money(mapped.get("price"))
    if sell_price is not None:
        payload["attributes"]["default_price_cents"] = sell_price
    return payload, errors


def preview_inventory_import(
    session: Session,
    workspace_id: uuid.UUID,
    rows: list[dict[str, Any]],
    mapping: dict[str, str],
    *,
    full_snapshot: bool = False,
    default_category: str = ItemCategory.GENERAL,
) -> dict[str, Any]:
    if "sku" not in mapping.values() or "title" not in mapping.values():
        raise ValueError("Map both a SKU and Title column before importing")

    existing = {
        row.sku: row
        for row in session.execute(
            select(models.InventoryItem).where(models.InventoryItem.workspace_id == workspace_id)
        ).scalars()
    }
    existing_by_isbn: dict[str, list[models.InventoryItem]] = {}
    for item in existing.values():
        isbn = str((item.attributes or {}).get("isbn") or "")
        if isbn:
            existing_by_isbn.setdefault(isbn, []).append(item)

    seen: set[str] = set()
    preview_rows: list[dict[str, Any]] = []
    counts = {"new": 0, "update": 0, "unchanged": 0, "conflict": 0}
    incoming_skus: set[str] = set()

    for index, raw in enumerate(rows, start=2):
        payload, errors = _row_payload(
            raw,
            mapping,
            default_category=default_category,
        )
        sku = payload.get("sku")
        action = "new"
        potential_duplicates: list[dict[str, str]] = []
        if sku:
            if sku in seen:
                errors.append("duplicate SKU in uploaded file")
            seen.add(sku)
            incoming_skus.add(sku)
        if errors:
            action = "conflict"
        elif sku in existing:
            item = existing[sku]
            comparable = {
                "title": item.title,
                "category": item.category,
                "quantity": item.quantity,
                "condition": item.condition,
                "cost_cents": item.cost_cents,
                "currency": item.currency,
                "location": item.location,
                "notes": item.notes,
                "attributes": item.attributes or {},
            }
            wanted = {key: payload[key] for key in comparable}
            action = "unchanged" if comparable == wanted else "update"
        else:
            isbn = str(payload.get("attributes", {}).get("isbn") or "")
            for item in existing_by_isbn.get(isbn, []) if isbn else []:
                potential_duplicates.append({"sku": item.sku, "title": item.title})

        if errors:
            action = "conflict"
        counts[action] += 1
        if len(preview_rows) < 200:
            preview_rows.append({
                "row": index,
                "action": action,
                "sku": sku,
                "title": payload.get("title"),
                "errors": errors,
                "potential_duplicates": potential_duplicates,
                "payload": payload,
            })

    missing_existing: list[dict[str, str]] = []
    if full_snapshot:
        for sku, item in existing.items():
            if item.status == ItemStatus.ACTIVE and sku not in incoming_skus:
                missing_existing.append({"sku": sku, "title": item.title})

    return {
        "rows": len(rows),
        "counts": counts,
        "preview": preview_rows,
        "missing_existing": missing_existing[:200],
        "missing_existing_count": len(missing_existing),
        "can_apply": counts["conflict"] == 0,
        "full_snapshot": full_snapshot,
        "default_category": default_category,
    }


def apply_inventory_import(
    session: Session,
    workspace_id: uuid.UUID,
    rows: list[dict[str, Any]],
    mapping: dict[str, str],
    *,
    full_snapshot: bool = False,
    default_category: str = ItemCategory.GENERAL,
) -> dict[str, Any]:
    preview = preview_inventory_import(
        session,
        workspace_id,
        rows,
        mapping,
        full_snapshot=full_snapshot,
        default_category=default_category,
    )
    if not preview["can_apply"]:
        raise ValueError("Import has conflicts. Resolve them before applying.")

    existing = {
        row.sku: row
        for row in session.execute(
            select(models.InventoryItem).where(models.InventoryItem.workspace_id == workspace_id)
        ).scalars()
    }
    incoming: set[str] = set()
    created = updated = unchanged = 0
    for raw in rows:
        payload, errors = _row_payload(
            raw,
            mapping,
            default_category=default_category,
        )
        if errors:
            raise ValueError("Import changed after preview and now contains invalid rows")
        sku = str(payload["sku"])
        incoming.add(sku)
        item = existing.get(sku)
        if item is None:
            item = models.InventoryItem(workspace_id=workspace_id, **payload)
            session.add(item)
            existing[sku] = item
            created += 1
            continue
        before = (
            item.title, item.category, item.quantity, item.condition, item.cost_cents,
            item.currency, item.location, item.notes, dict(item.attributes or {})
        )
        item.title = payload["title"]
        item.category = payload["category"]
        item.quantity = int(payload["quantity"])
        item.condition = payload["condition"]
        item.cost_cents = payload["cost_cents"]
        item.currency = payload["currency"]
        item.location = payload["location"]
        item.notes = payload["notes"]
        item.attributes = payload["attributes"]
        item.status = ItemStatus.ACTIVE if item.quantity > 0 else ItemStatus.ARCHIVED
        after = (
            item.title, item.category, item.quantity, item.condition, item.cost_cents,
            item.currency, item.location, item.notes, dict(item.attributes or {})
        )
        if before == after:
            unchanged += 1
        else:
            updated += 1

    archived = 0
    if full_snapshot:
        for sku, item in existing.items():
            if sku not in incoming and item.status == ItemStatus.ACTIVE:
                item.status = ItemStatus.ARCHIVED
                item.quantity = 0
                archived += 1

    session.flush()
    return {
        "created": created,
        "updated": updated,
        "unchanged": unchanged,
        "archived": archived,
        "rows": len(rows),
    }


EXPORT_HEADERS = [
    "SKU", "Title", "Category", "Quantity", "Condition", "Cost", "Price",
    "Currency", "Location", "Notes", "Barcode", "Author", "ISBN", "Subtitle",
    "Publisher", "Edition", "Binding", "Language", "Publish Date",
    "Publication Year", "Pages", "Brand", "Size", "Colour", "Material",
    "Measurements", "Status",
]


def inventory_export_rows(items: Iterable[models.InventoryItem]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in items:
        attrs = item.attributes or {}
        rows.append({
            "SKU": item.sku,
            "Title": item.title,
            "Category": item.category,
            "Quantity": item.quantity,
            "Condition": item.condition or "",
            "Cost": f"{item.cost_cents / 100:.2f}" if item.cost_cents is not None else "",
            "Price": (
                f"{int(attrs['default_price_cents']) / 100:.2f}"
                if attrs.get("default_price_cents") is not None else ""
            ),
            "Currency": item.currency or "EUR",
            "Location": item.location or "",
            "Notes": item.notes or "",
            "Barcode": attrs.get("barcode", ""),
            "Author": attrs.get("author", ""),
            "ISBN": attrs.get("isbn", ""),
            "Subtitle": attrs.get("subtitle", ""),
            "Publisher": attrs.get("publisher", ""),
            "Edition": attrs.get("edition", ""),
            "Binding": attrs.get("binding", ""),
            "Language": attrs.get("language", ""),
            "Publish Date": attrs.get("publish_date", ""),
            "Publication Year": attrs.get("publication_year", ""),
            "Pages": attrs.get("pages", ""),
            "Brand": attrs.get("brand", ""),
            "Size": attrs.get("size", ""),
            "Colour": attrs.get("colour", ""),
            "Material": attrs.get("material", ""),
            "Measurements": attrs.get("measurements", ""),
            "Status": item.status,
        })
    return rows


def render_csv(
    rows: list[dict[str, Any]],
    headers: list[str] | None = None,
) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=headers or EXPORT_HEADERS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def render_xlsx(
    rows: list[dict[str, Any]],
    headers: list[str] | None = None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory"
    selected_headers = headers or EXPORT_HEADERS
    sheet.append(selected_headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in selected_headers])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
